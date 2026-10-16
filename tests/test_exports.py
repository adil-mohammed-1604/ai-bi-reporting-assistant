import io

import pandas as pd

from src.export_utils import to_safe_csv


def roundtrip(df):
    return pd.read_csv(io.BytesIO(to_safe_csv(df)), dtype=str, keep_default_na=False)


def test_formula_like_text_is_neutralised():
    df = pd.DataFrame({"Description": ['=HYPERLINK("http://x","y")', "+cmd|calc", "@SUM(A1)", "-evil", "MUG"]})
    out = roundtrip(df)["Description"].tolist()
    assert out[0].startswith("'=") and out[1].startswith("'+") and out[2].startswith("'@")
    assert out[3].startswith("'-")
    assert out[4] == "MUG"


def test_plain_numbers_and_negatives_are_untouched():
    df = pd.DataFrame({"Quantity": ["-5", "+3", "2.5", "7"]})
    assert roundtrip(df)["Quantity"].tolist() == ["-5", "+3", "2.5", "7"]


def test_non_text_columns_and_missing_values_survive():
    df = pd.DataFrame({"n": [1, 2], "t": ["a", None]})
    out = roundtrip(df)
    assert out["n"].tolist() == ["1", "2"] and out["t"].tolist() == ["a", ""]


def test_csv_has_utf8_bom_for_excel():
    assert to_safe_csv(pd.DataFrame({"a": ["é"]})).startswith(b"\xef\xbb\xbf")


def test_original_dataframe_not_modified():
    df = pd.DataFrame({"t": ["=1+1x"]})
    to_safe_csv(df)
    assert df["t"].iloc[0] == "=1+1x"


# ------------------------------------------------------------------ Excel workbook
import io as _io

import openpyxl as _openpyxl

from src.export_utils import build_kpi_workbook
from src.ui_helpers import SAMPLE_PATH, process


def _sample_book(**kw):
    up = process(SAMPLE_PATH.read_bytes(), SAMPLE_PATH.name)
    raw = build_kpi_workbook(up["name"], up["hash"], up["quality"], up["kpis"], **kw)
    return up, _openpyxl.load_workbook(_io.BytesIO(raw))


def test_workbook_has_expected_sheets_and_numbers():
    up, wb = _sample_book(run_id=2, reconciled=True)
    assert wb.sheetnames == ["Summary", "Monthly", "Countries", "Top Products",
                             "Data Quality", "Rejected Rows", "Flagged Rows"]
    cells = {c.value for row in wb["Summary"].iter_rows() for c in row}
    assert "Agreed" in cells and 346087.77 in cells
    assert wb["Rejected Rows"].max_row == 1 + 9
    assert wb["Flagged Rows"].max_row == 1 + 149


def test_workbook_marks_disagreement_and_unsaved():
    _, wb = _sample_book(run_id=3, reconciled=False)
    assert any("DISAGREED" in str(c.value) for row in wb["Summary"].iter_rows() for c in row)
    _, wb = _sample_book()
    assert any(c.value == "not saved" for row in wb["Summary"].iter_rows() for c in row)


def test_workbook_never_stores_uploaded_text_as_formula():
    csv = (b"InvoiceNo,StockCode,Quantity,InvoiceDate,UnitPrice,Description\n"
           b"A1,P1,2,2025-01-05,5,=HYPERLINK(\"http://evil\")\n"
           b"A2,P2,1,not a date,5,x\n")
    up = process(csv, "evil.csv")
    wb = _openpyxl.load_workbook(_io.BytesIO(
        build_kpi_workbook(up["name"], up["hash"], up["quality"], up["kpis"])))
    for ws in wb:
        for row in ws.iter_rows():
            for c in row:
                assert c.data_type != "f", (ws.title, c.coordinate)