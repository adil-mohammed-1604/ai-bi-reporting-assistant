"""Export helpers.

Uploaded files are untrusted data. A cell such as =HYPERLINK(...) written back
out to CSV can run as a formula when someone opens the export in Excel, so
text cells that could be read as formulas are neutralised with a leading
apostrophe. Plain numbers such as -5 are left alone.
"""
import pandas as pd

_FORMULA_TRIGGERS = ("=", "+", "-", "@", "\t", "\r")


def _neutralise(value):
    if isinstance(value, str) and value and value[0] in _FORMULA_TRIGGERS:
        try:
            float(value)
            return value          # an ordinary number like -5 or +3.2
        except ValueError:
            return "'" + value
    return value


def to_safe_csv(df: pd.DataFrame) -> bytes:
    """CSV bytes (UTF-8 with BOM so Excel reads accents correctly), formula-safe."""
    out = df.copy()
    for col in out.columns:
        if pd.api.types.is_string_dtype(out[col]) or pd.api.types.is_object_dtype(out[col]):
            out[col] = out[col].map(_neutralise)
    return out.to_csv(index=False).encode("utf-8-sig")


# ---------------------------------------------------------------- Excel workbook
import io
import re
from datetime import datetime, timezone

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from src import config

_ILLEGAL_XML = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")
_HEADER_FILL = PatternFill("solid", fgColor="DCE6F2")
MAX_ROW_SHEET = 50_000       # keep the sheets for flagged/rejected rows a sensible size


def _plain(value):
    """Convert numpy/pandas values to plain Python; missing values become None."""
    if value is None or (not isinstance(value, str) and pd.isna(value)):
        return None
    if hasattr(value, "item") and not isinstance(value, (str, bytes)):
        value = value.item()
    if isinstance(value, pd.Timestamp):
        return value.to_pydatetime()
    return value


def _put(ws, row, col, value, number_format=None, bold=False, wrap=False):
    """Write one cell. Text is always stored as text, so a value such as
    =HYPERLINK(...) in an uploaded file can never become a live formula."""
    value = _plain(value)
    cell = ws.cell(row=row, column=col)
    if isinstance(value, str):
        cell.value = _ILLEGAL_XML.sub("", value)
        cell.data_type = "s"
    else:
        cell.value = value
    if number_format:
        cell.number_format = number_format
    if bold:
        cell.font = Font(bold=True)
    if wrap:
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    return cell


def _table(ws, start_row, headers, rows, formats=None, wrap_cols=()):
    formats = formats or {}
    for j, h in enumerate(headers, 1):
        c = _put(ws, start_row, j, h, bold=True)
        c.fill = _HEADER_FILL
    r = start_row
    for r, row in enumerate(rows, start_row + 1):
        for j, v in enumerate(row, 1):
            _put(ws, r, j, v, number_format=formats.get(headers[j - 1]), wrap=headers[j - 1] in wrap_cols)
    return r + 1


def _widths(ws, widths):
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def build_kpi_workbook(file_name, file_hash, quality, kpis, run_id=None, reconciled=None, report=None) -> bytes:
    """Excel workbook with the KPI summary, definitions, chart tables and quality results."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    s = quality.summary

    _put(ws, 1, 1, "Business KPI Summary", bold=True).font = Font(bold=True, size=14)
    info = [
        ("File", file_name),
        ("Period analysed", kpis["period"]),
        ("Generated", datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M UTC")),
        ("File fingerprint (SHA-256, first 16)", (file_hash or "")[:16]),
        ("SQL Server run ID", run_id if run_id is not None else "not saved"),
        ("Python vs SQL reconciliation",
         {True: "Agreed", False: "DISAGREED: treat figures as unverified", None: "Not run"}[reconciled]),
        ("Rows in file / eligible for KPIs", f"{s['original_rows']:,} / {s['eligible_rows']:,}"),
        ("Rows rejected / flagged", f"{s['rejected_rows']:,} / {s['flagged_rows']:,}"),
    ]
    for i, (k, v) in enumerate(info, 3):
        _put(ws, i, 1, k, bold=True)
        _put(ws, i, 2, v)
    start = len(info) + 5
    rows = [(m["name"], m["value"], m["period"], m["definition"], m["limitations"]) for m in kpis["metrics"]]
    _table(ws, start, ["Metric", "Value", "Period", "Definition", "Limitations"], rows,
           formats={"Value": "#,##0.00"}, wrap_cols=("Definition", "Limitations"))
    _widths(ws, [38, 20, 24, 60, 70])
    ws.freeze_panes = ws.cell(row=start + 1, column=1)

    ws = wb.create_sheet("Monthly")
    m = kpis["monthly"]
    _table(ws, 1, ["Month", "Revenue", "Orders"], m.itertuples(index=False),
           formats={"Revenue": "#,##0.00", "Orders": "#,##0"})
    _widths(ws, [12, 16, 10]); ws.freeze_panes = "A2"

    ws = wb.create_sheet("Countries")
    _table(ws, 1, ["Country", "Revenue", "Orders"], kpis["countries"].itertuples(index=False),
           formats={"Revenue": "#,##0.00", "Orders": "#,##0"})
    _widths(ws, [22, 16, 10]); ws.freeze_panes = "A2"

    ws = wb.create_sheet("Top Products")
    p = kpis["products"][["StockCode", "Description", "Revenue", "UnitsSold"]]
    _table(ws, 1, ["StockCode", "Description", "Revenue", "UnitsSold"], p.itertuples(index=False),
           formats={"Revenue": "#,##0.00", "UnitsSold": "#,##0"})
    _widths(ws, [12, 36, 16, 12]); ws.freeze_panes = "A2"

    ws = wb.create_sheet("Data Quality")
    summary_rows = [("Original rows", s["original_rows"]), ("Accepted", s["accepted_rows"]),
                    ("Flagged", s["flagged_rows"]), ("Rejected", s["rejected_rows"]),
                    ("Eligible for KPIs", s["eligible_rows"]), ("Acceptance rate %", s["acceptance_rate"])]
    nxt = _table(ws, 1, ["Measure", "Value"], summary_rows)
    _table(ws, nxt + 1, ["ColumnName", "IssueType", "Severity", "IssueCount", "Details"],
           quality.issues.itertuples(index=False), wrap_cols=("Details",))
    _widths(ws, [24, 26, 12, 12, 80])

    cols = ["SourceRow"] + config.ALL_COLUMNS + ["Status", "Reasons"]
    for title, frame in (("Rejected Rows", quality.rejected), ("Flagged Rows", quality.flagged)):
        ws = wb.create_sheet(title)
        shown = frame[cols].head(MAX_ROW_SHEET)
        end = _table(ws, 1, cols, shown.itertuples(index=False))
        if len(frame) > MAX_ROW_SHEET:
            _put(ws, end + 1, 1, f"Only the first {MAX_ROW_SHEET:,} of {len(frame):,} rows are shown here.")
        _widths(ws, [10] + [16] * len(config.ALL_COLUMNS) + [12, 40]); ws.freeze_panes = "A2"

    if report is not None and report.status == "VERIFIED":
        ws = wb.create_sheet("Report Draft")
        c = report.content
        lines = [("DRAFT: requires human review", ""), ("Source", report.provider + (f" / {report.model}" if report.model else "")),
                 ("Generated", report.generated_at), ("Executive summary", c["executive_summary"])]
        lines += [("Key finding", t) for t in c["key_findings"]]
        lines += [("Follow-up analysis", t) for t in c["follow_up_analysis"]]
        lines += [("Limitation", t) for t in c["limitations"]]
        for i, (k, v) in enumerate(lines, 1):
            _put(ws, i, 1, k, bold=True)
            _put(ws, i, 2, v, wrap=True)
        _widths(ws, [26, 110])

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()